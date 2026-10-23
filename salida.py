"""Escribe el Excel clasificado de una empresa para un mes."""
import datetime as dt
import pandas as pd
from comun import SECCIONES_POR_OPERACION, tipo_operacion
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AMARILLO = PatternFill('solid', fgColor='FFF2CC')
GRIS = PatternFill('solid', fgColor='EDEDED')
AZUL = PatternFill('solid', fgColor='DDEBF7')
ROJO = PatternFill('solid', fgColor='F8CBAD')
B = Font(bold=True)
NUM = '#,##0.00;[Red]-#,##0.00'
LINEA = Border(top=Side(style='thin'))
MESES = ['', 'ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO', 'AGOSTO', 'SETIEMBRE', 'OCTUBRE',
         'NOVIEMBRE', 'DICIEMBRE']


def secciones(mov, modelo, empresa):
    """Lista ordenada de (tipo, categoria): las que la empresa suele usar + las que aparecen este mes."""
    usadas = set()
    for tipo in ('INGRESO', 'EGRESO'):
        for c in modelo.permitidas.get((empresa, tipo), set()):
            usadas.add((tipo, c))
    usadas |= set(zip(mov.tipo_seccion, mov.categoria))
    return sorted(usadas, key=lambda tc: (tc[0] != 'INGRESO', modelo.orden_de(empresa, tc[0], tc[1]), tc[1]))


def escribir(ruta, empresa, anio, mes, mov, cuentas, modelo, umbral):
    mov = mov.copy()
    # el ITF que la empresa deja dentro del movimiento va en la misma sección que la otra parte de la fila
    mov['tipo_seccion'] = mov.tipo
    if 'itf_misma_fila' in mov:
        for i in mov.index[(mov.itf_misma_fila == 1) & (mov.categoria != 'ITF')]:
            otra = mov[(mov.hoja == mov.hoja[i]) & (mov.fila_origen == mov.fila_origen[i]) & (mov.index != i)]
            if len(otra):
                mov.loc[i, 'tipo_seccion'] = otra.tipo.iloc[0]
    mov['revisar'] = (mov.confianza < umbral)
    secs = secciones(mov, modelo, empresa)
    wb = Workbook()

    # ---------- RESUMEN ----------
    ws = wb.active
    ws.title = 'RESUMEN'
    ws['A1'] = f'{empresa} – FLUJO {MESES[mes]} {anio} (clasificación automática)'
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = (f'Generado {dt.datetime.now():%d/%m/%Y %H:%M}. Filas amarillas = revisar (confianza < {umbral:.0%}). '
                'Si corriges una categoría en MOVIMIENTOS, este resumen se recalcula solo; '
                'las hojas por cuenta se rehacen con actualizar.py.')
    ws['A2'].font = Font(italic=True, color='666666')
    hdr = ['CUENTA', 'Nº CUENTA', 'MONEDA', 'SALDO INICIAL', 'INGRESOS', 'EGRESOS', 'SALDO FINAL (calc.)',
           'SALDO SEGÚN LIBRO', 'DIFERENCIA', 'MOVIMIENTOS', 'A REVISAR']
    ws.append([])
    ws.append(hdr)
    for c in ws[4]:
        c.font, c.fill = B, AZUL
    cu = cuentas[cuentas.movimientos > 0].reset_index(drop=True)
    for k, r in cu.iterrows():
        f = 5 + k
        ws.append([r.hoja, r.cuenta, r.moneda, r.saldo_inicial, r.ingresos, r.egresos, f'=D{f}+E{f}-F{f}',
                   r.saldo_final, f'=ROUND(G{f}-H{f},2)', r.movimientos, int(mov[mov.hoja == r.hoja].revisar.sum())])
        for col in 'DEFGHI':
            ws[f'{col}{f}'].number_format = NUM
        if abs(r.saldo_inicial + r.ingresos - r.egresos - r.saldo_final) > 0.05:
            for col in 'GHI':
                ws[f'{col}{f}'].fill = ROJO
    fin = 5 + len(cu)
    ws.cell(fin, 1, 'Rojo = el saldo del libro no coincide con saldo inicial + movimientos del mes '
                    '(suele ser una fila con fecha de otro mes dentro del bloque, o saldo mal arrastrado).').font = \
        Font(italic=True, color='C00000')
    sin = cuentas[cuentas.movimientos == 0].hoja.tolist()
    if sin:
        ws.cell(fin + 1, 1, 'Cuentas sin movimientos en el mes: ' + ', '.join(sin)).font = Font(italic=True, color='666666')

    # matriz categoría x cuenta con SUMIFS (se recalcula si corrigen MOVIMIENTOS)
    f0 = fin + 3
    ws.cell(f0, 1, 'TOTALES POR CATEGORÍA (en la moneda de cada cuenta)').font = B
    ws.cell(f0 + 1, 1, 'TIPO').font = B
    ws.cell(f0 + 1, 2, 'CATEGORÍA').font = B
    for j, h in enumerate(cu.hoja):
        c = ws.cell(f0 + 1, 3 + j, f'{h} ({cu.moneda[j]})')
        c.font, c.fill = B, AZUL
    n = len(mov) + 1
    for i, (tipo, cat) in enumerate(secs):
        f = f0 + 2 + i
        ws.cell(f, 1, tipo)
        ws.cell(f, 2, cat)
        for j, h in enumerate(cu.hoja):
            col = get_column_letter(3 + j)
            c = ws.cell(f, 3 + j, f'=SUMIFS(MOVIMIENTOS!$F$2:$F${n},MOVIMIENTOS!$A$2:$A${n},"{h}",'
                                  f'MOVIMIENTOS!$G$2:$G${n},$B{f},MOVIMIENTOS!$E$2:$E${n},$A{f})')
            c.number_format = NUM
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 34
    for j in range(3, 3 + max(len(cu), 9)):
        ws.column_dimensions[get_column_letter(j)].width = 16

    # ---------- MOVIMIENTOS (editable) ----------
    wm = wb.create_sheet('MOVIMIENTOS')
    cols = ['CUENTA', 'MONEDA', 'FECHA', 'CONCEPTO', 'TIPO', 'MONTO', 'CATEGORIA', 'CONFIANZA', 'ALTERNATIVA',
            'REVISAR', 'ORIGEN', 'NOTA', 'Nº OPERACION', 'CH/Nº', 'FILA EN LIBRO']
    wm.append(cols)
    for c in wm[1]:
        c.font, c.fill = B, AZUL
    orden_hojas = {h: i for i, h in enumerate(cuentas.hoja)}
    mv = mov.assign(_o=mov.hoja.map(orden_hojas)).sort_values(['_o', 'fecha', 'fila_origen'])
    for _, r in mv.iterrows():
        nota = r.get('aviso', '') or ''
        if r.get('itf_misma_fila', 0) == 1:
            nota = (nota + ' ITF que venía en la misma fila de: ' + str(r.get('concepto_padre', ''))).strip()
        wm.append([r.hoja, r.moneda, r.fecha, r.concepto, r.tipo, r.monto, r.categoria, r.confianza,
                   r.alternativa, 'SI' if r.revisar else '', r.origen, nota,
                   r.operacion if pd.notna(r.operacion) else None, r.cheque if pd.notna(r.cheque) else None,
                   int(r.fila_origen)])
        if r.revisar:
            for c in wm[wm.max_row]:
                c.fill = AMARILLO
    for row in wm.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = 'DD/MM/YYYY'
    for row in wm.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].number_format = NUM
    for row in wm.iter_rows(min_row=2, min_col=8, max_col=8):
        row[0].number_format = '0%'
    for col, w in zip('ABCDEFGHIJKLMNO', [16, 7, 11, 42, 9, 14, 30, 10, 28, 8, 10, 40, 14, 8, 8]):
        wm.column_dimensions[col].width = w
    wm.freeze_panes = 'A2'
    wm.auto_filter.ref = f'A1:O{wm.max_row}'
    # lista desplegable de categorías
    wl = wb.create_sheet('LISTAS')
    cats = sorted({c for _, c in secs})
    for i, c in enumerate(cats, 1):
        wl.cell(i, 1, c)
    wl.sheet_state = 'hidden'
    dv = DataValidation(type='list', formula1=f'=LISTAS!$A$1:$A${len(cats)}', allow_blank=False,
                        showErrorMessage=False)
    wm.add_data_validation(dv)
    dv.add(f'G2:G{max(wm.max_row, 2)}')

    # ---------- una hoja por cuenta, agrupada como sus RESUMEN FLUJO ----------
    for _, cta in cuentas.iterrows():
        m = mov[mov.hoja == cta.hoja]
        if m.empty:
            continue
        wc = wb.create_sheet(cta.hoja[:31])
        wc['A1'] = empresa
        wc['A2'] = cta.cuenta or cta.hoja
        wc['A3'] = f'{MESES[mes]} {anio} – {cta.moneda}'
        for a in ('A1', 'A2', 'A3'):
            wc[a].font = B
        wc['C5'], wc['C6'] = 'SALDO INICIAL', 'SALDO FINAL'
        wc['D5'] = cta.saldo_inicial
        wc['A7'], wc['B7'], wc['C7'], wc['D7'], wc['F7'], wc['G7'] = \
            'FECHA', 'CONCEPTO', 'DEBE', 'HABER', 'CONF.', 'Nº OPERACION'
        for c in wc[7]:
            c.font, c.fill = B, AZUL
        f = 9
        filas_debe, filas_haber = [], []
        for tipo_sec, cat in secs:
            s = m[(m.categoria == cat) & (m.tipo_seccion == tipo_sec)]
            if cat in SECCIONES_POR_OPERACION:  # overnight, BCR, cambio de moneda...: por tipo y luego por fecha
                s = s.assign(_op=s.concepto.map(tipo_operacion)).sort_values(['_op', 'fecha', 'fila_origen'])
            else:
                s = s.sort_values(['fecha', 'fila_origen'])
            wc.cell(f, 1, cat).font = B
            wc.cell(f, 1).fill = GRIS
            f += 1
            ini = f
            for _, r in s.iterrows():
                wc.cell(f, 1, r.fecha).number_format = 'DD/MM/YYYY'
                wc.cell(f, 2, r.concepto)
                wc.cell(f, 3 if r.tipo == 'INGRESO' else 4, r.monto).number_format = NUM
                wc.cell(f, 6, r.confianza).number_format = '0%'
                if pd.notna(r.operacion):
                    wc.cell(f, 7, r.operacion)
                if r.revisar:
                    for col in range(1, 8):
                        wc.cell(f, col).fill = AMARILLO
                f += 1
            colsub = 'C' if tipo_sec == 'INGRESO' else 'D'
            c = wc.cell(f, 5, f'=SUM(C{ini}:D{f - 1})' if f > ini else 0)
            c.number_format, c.font, c.border = NUM, B, LINEA
            if f > ini:
                filas_debe.append(f'C{ini}:C{f - 1}')
                filas_haber.append(f'D{ini}:D{f - 1}')
            f += 2
        wc.cell(f, 2, 'TOTALES').font = B
        wc.cell(f, 3, '=' + ('+'.join(f'SUM({x})' for x in filas_debe) or '0')).number_format = NUM
        wc.cell(f, 4, '=' + ('+'.join(f'SUM({x})' for x in filas_haber) or '0')).number_format = NUM
        wc.cell(f, 3).font = wc.cell(f, 4).font = B
        wc['D6'] = f'=D5+C{f}-D{f}'
        wc['E6'] = cta.saldo_final
        wc['F6'] = '=ROUND(D6-E6,2)'
        wc['E5'] = 'según libro →'
        wc['E5'].font = Font(italic=True, color='666666')
        for a in ('D5', 'D6', 'E6', 'F6'):
            wc[a].number_format = NUM
        wc.column_dimensions['A'].width = 12
        wc.column_dimensions['B'].width = 44
        for col in 'CDE':
            wc.column_dimensions[col].width = 15
        wc.column_dimensions['G'].width = 14
        wc.freeze_panes = 'A8'
    wb.move_sheet('MOVIMIENTOS', offset=-(len(wb.sheetnames) - 2))
    wb.save(ruta)
    return mov
